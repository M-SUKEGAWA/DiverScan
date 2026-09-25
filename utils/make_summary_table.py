"""
SummaryTable.xlsx generator

✅ Supports:
- model_type: "multi" or "single"
- class_name_list: variable number of conditions (sheet names must match)
- number_of_attention_branches: variable number of attentions

Behavior:
- For correlation-like tables (corr): output top3 + bottom3 (or all if <= 6).
- For edge tables: output all rows (sorted desc, NaN last).
- For cluster tables: output all rows (sorted desc, NaN last). Feature cast to int if numeric.
- For Hellinger tables: detect columns "Attention{a} Condition i - j".
  For each pair-column: sort desc (NaN last), output top3 (or all if <= 4).
- Node attention sheet is created only for model_type == "multi".
"""

from __future__ import annotations

import argparse
import ast
from pathlib import Path
from typing import List, Sequence, Tuple

import numpy as np
import pandas as pd
import openpyxl


# --------------------------
# Sections
# --------------------------
# Section titles are kept identical to the feature-group names shown in the notebook
# (the "show" dropdown labels), so the SummaryTable rows match what the user sees there.
SECTIONS_MULTI = [
    ("Correlation (basic features)", "correlation.xlsx", "corr"),
    ("Correlation in highlighted segments (basic features)", "correlation_in_highlighted_segments.xlsx", "corr"),
    ("Hellinger distance in highlighted segments (basic features)", "hellinger_distance.xlsx", "hell"),

    ("[Multi mode] Correlation (differences of basic features between nodes)", "correlation_node_diff.xlsx", "corr"),
    ("[Multi mode] Correlation in highlighted segments (differences of basic features between nodes)",
     "correlation_node_diff_in_highlighted_segments.xlsx", "corr"),
    ("[Multi mode] Hellinger distance in highlighted segments (differences of basic features between nodes)",
     "hellinger_distance_node_diff.xlsx", "hell"),

    ("[Multi mode] Correlation (edge features)", "correlation_edge.xlsx", "edge"),
    ("[Multi mode] Correlation in highlighted segments (edge features)", "correlation_edge_in_highlighted_segments.xlsx", "edge"),
    ("[Multi mode] Hellinger distance in highlighted segments (edge features)", "hellinger_distance_edge.xlsx", "hell"),

    ("[Multi mode] Correlation (incident edge features per node)", "correlation_edge_per_node.xlsx", "corr"),
    ("[Multi mode] Correlation in highlighted segments (incident edge features per node)",
     "correlation_edge_per_node_in_highlighted_segments.xlsx", "corr"),
    ("[Multi mode] Hellinger distance in highlighted segments (incident edge features per node)",
     "hellinger_distance_edge_per_node.xlsx", "hell"),

    ("[Multi mode] Correlation (MST features)", "correlation_mst.xlsx", "corr"),
    ("[Multi mode] Correlation in highlighted segments (MST features)", "correlation_mst_in_highlighted_segments.xlsx", "corr"),
    ("[Multi mode] Hellinger distance in highlighted segments (MST features)", "hellinger_distance_mst.xlsx", "hell"),

    ("[Multi mode] Correlation (geometric features)", "correlation_geometric.xlsx", "corr"),
    ("[Multi mode] Correlation in highlighted segments (geometric features)",
     "correlation_geometric_in_highlighted_segments.xlsx", "corr"),
    ("[Multi mode] Hellinger distance in highlighted segments (geometric features)",
     "hellinger_distance_geometric.xlsx", "hell"),

    ("[Multi mode] Correlation (cluster number based on the distance between nodes)", "correlation_cluster.xlsx", "cluster"),
    ("[Multi mode] Correlation in highlighted segments (cluster number based on the distance between nodes)",
     "correlation_cluster_in_highlighted_segments.xlsx", "cluster"),

    ("[Multi mode] Correlation (configuration patterns based on the distance between nodes)", "correlation_config.xlsx", "cluster"),
    ("[Multi mode] Correlation in highlighted segments (configuration patterns based on the distance between nodes)",
     "correlation_config_in_highlighted_segments.xlsx", "cluster"),
]

SECTIONS_SINGLE = [
    ("Correlation (basic features)", "correlation.xlsx", "corr"),
    ("Correlation in highlighted segments (basic features)", "correlation_in_highlighted_segments.xlsx", "corr"),
    ("Hellinger distance in highlighted segments (basic features)", "hellinger_distance.xlsx", "hell"),
]


# --------------------------
# Parsing / Value conversion
# --------------------------
def parse_class_names(arg: str) -> List[str]:
    """
    Accept:
      - Python list literal string: "['a','b']" or '["a","b"]'
      - Comma-separated: "a,b"
    """
    s = arg.strip()
    if not s:
        return []

    if s.startswith("[") and s.endswith("]"):
        try:
            v = ast.literal_eval(s)
        except Exception as e:
            raise ValueError(f"--class-names list parse failed: {e}\nGiven: {arg}") from e
        if not isinstance(v, (list, tuple)):
            raise ValueError(f"--class-names must be a list/tuple. Given: {type(v)}")
        out: List[str] = []
        for x in v:
            if not isinstance(x, str):
                raise ValueError(f"--class-names elements must be strings. Found: {type(x)} ({x})")
            x2 = x.strip()
            if x2:
                out.append(x2)
        return out

    return [x.strip() for x in s.split(",") if x.strip()]


def to_excel_value(x):
    """
    Convert to Excel cell value:
      - NaN/NA -> "nan" (string)
      - float -> rounded to 6 decimals
      - int -> int
      - otherwise -> as-is
    """
    if x is None:
        return None
    try:
        if pd.isna(x):
            return "nan"
    except Exception:
        pass

    if isinstance(x, (float, np.floating)):
        return float(np.round(float(x), 6))
    if isinstance(x, (int, np.integer)):
        return int(x)
    return x


# --------------------------
# Data helpers
# --------------------------
def stable_sort_desc(df: pd.DataFrame, col: str) -> pd.DataFrame:
    """Stable sort descending, NaNs last."""
    return df.sort_values(col, ascending=False, na_position="last", kind="mergesort").reset_index(drop=True)


def pick_top_bottom_indices(n: int, top: int = 3, bottom: int = 3) -> List[int]:
    """Return indices for top 'top' and bottom 'bottom' rows (no duplicates)."""
    if n <= top + bottom:
        return list(range(n))
    idx = list(range(top)) + list(range(n - bottom, n))
    out, seen = [], set()
    for i in idx:
        if i not in seen:
            out.append(i)
            seen.add(i)
    return out


def validate_sheets_exist(path: Path, class_name_list: Sequence[str]) -> None:
    xl = pd.ExcelFile(path)
    missing = [c for c in class_name_list if c not in xl.sheet_names]
    if missing:
        raise ValueError(
            f"Missing sheets in {path.name}: {missing}\n"
            f"Available sheets: {xl.sheet_names}"
        )


def parse_hellinger_pair_columns(columns: Sequence[str], att: int) -> List[Tuple[str, int, int]]:
    """
    Find columns like: "Attention{att} Condition i - j"
    Return list of tuples: (col_name, i, j)
    """
    prefix = f"Attention{att} Condition "
    out: List[Tuple[str, int, int]] = []
    for c in columns:
        if not isinstance(c, str) or not c.startswith(prefix):
            continue
        tail = c[len(prefix):].strip()
        parts = [p.strip() for p in tail.split("-")]
        if len(parts) != 2:
            continue
        try:
            i = int(parts[0])
            j = int(parts[1])
        except ValueError:
            continue
        out.append((c, i, j))
    out.sort(key=lambda x: (x[1], x[2]))
    return out


# --------------------------
# Row builders
# --------------------------
def build_corr_rows(
    input_dir: Path, fname: str, att: int, title: str, class_name_list: Sequence[str]
) -> List[List[object]]:
    """
    Correlation-like:
      - EXCLUDE NaN rows from ranking (drop NaN in the attention column)
      - stable sort desc
      - output top3 + bottom3 (or all if <= 6)
    """
    rows: List[List[object]] = []
    first_row_of_section = True
    col = f"Attention{att}"

    fpath = input_dir / fname
    validate_sheets_exist(fpath, class_name_list)

    for cond in class_name_list:
        df = pd.read_excel(fpath, sheet_name=cond)
        if "Feature" not in df.columns or col not in df.columns:
            raise ValueError(f"{fname}:{cond} must contain columns ['Feature', '{col}'].")

        df = df[["Feature", col]].dropna(subset=[col])  # ✅ exclude NaN from ranking
        df = stable_sort_desc(df, col)

        pick = pick_top_bottom_indices(len(df), top=3, bottom=3)

        for j, i in enumerate(pick):
            row = [None, None, i + 1, to_excel_value(df.loc[i, "Feature"]), to_excel_value(df.loc[i, col])]
            if first_row_of_section:
                row[0] = title
                first_row_of_section = False
            if j == 0:
                row[1] = cond
            rows.append(row)

    return rows


def build_edge_rows(
    input_dir: Path, fname: str, att: int, title: str, class_name_list: Sequence[str]
) -> List[List[object]]:
    """
    Edge correlation:
      - output ALL rows
      - NaN values are written as "nan"
    """
    rows: List[List[object]] = []
    first_row_of_section = True
    col = f"Attention{att}"

    fpath = input_dir / fname
    validate_sheets_exist(fpath, class_name_list)

    for cond in class_name_list:
        df = pd.read_excel(fpath, sheet_name=cond)
        if "Feature" not in df.columns or col not in df.columns:
            raise ValueError(f"{fname}:{cond} must contain columns ['Feature', '{col}'].")

        df = df[["Feature", col]]
        df = stable_sort_desc(df, col)

        for j in range(len(df)):
            row = [None, None, j + 1, to_excel_value(df.loc[j, "Feature"]), to_excel_value(df.loc[j, col])]
            if first_row_of_section:
                row[0] = title
                first_row_of_section = False
            if j == 0:
                row[1] = cond
            rows.append(row)

    return rows


def build_cluster_rows(
    input_dir: Path, fname: str, att: int, title: str, class_name_list: Sequence[str]
) -> List[List[object]]:
    """
    Cluster correlation:
      - output ALL rows
      - Feature cast to int if numeric
      - NaN values are written as "nan"
    """
    rows: List[List[object]] = []
    first_row_of_section = True
    col = f"Attention{att}"

    fpath = input_dir / fname
    validate_sheets_exist(fpath, class_name_list)

    for cond in class_name_list:
        df = pd.read_excel(fpath, sheet_name=cond)
        if "Feature" not in df.columns or col not in df.columns:
            raise ValueError(f"{fname}:{cond} must contain columns ['Feature', '{col}'].")

        df = df[["Feature", col]]
        df = stable_sort_desc(df, col)

        for j in range(len(df)):
            feat_raw = df.loc[j, "Feature"]
            if pd.api.types.is_number(feat_raw) and not pd.isna(feat_raw):
                feat = int(feat_raw)
            else:
                feat = to_excel_value(feat_raw)

            row = [None, None, j + 1, feat, to_excel_value(df.loc[j, col])]
            if first_row_of_section:
                row[0] = title
                first_row_of_section = False
            if j == 0:
                row[1] = cond
            rows.append(row)

    return rows


def build_hellinger_rows(
    input_dir: Path, fname: str, att: int, title: str, class_name_list: Sequence[str]
) -> List[List[object]]:
    """
    Hellinger:
      - detect columns "Attention{att} Condition i - j"
      - EXCLUDE NaN rows from ranking (drop NaN in that pair-column)
      - sort desc
      - output top3 (or all if <=4)
    """
    fpath = input_dir / fname
    df = pd.read_excel(fpath)

    pair_cols = parse_hellinger_pair_columns(df.columns, att)
    if not pair_cols:
        raise ValueError(
            f"No Hellinger pair columns found in {fname} for Attention{att}.\n"
            f"Expected columns like 'Attention{att} Condition 1 - 2'.\n"
            f"Found columns: {list(df.columns)}"
        )

    rows: List[List[object]] = []
    first_row_of_section = True

    for (col, i, j) in pair_cols:
        if 1 <= i <= len(class_name_list) and 1 <= j <= len(class_name_list):
            pair_label = f"{class_name_list[i-1]} - {class_name_list[j-1]}"
        else:
            pair_label = f"Condition {i} - {j}"

        sub = df[["Feature", col]].dropna(subset=[col])  # ✅ exclude NaN from ranking
        sub = stable_sort_desc(sub, col)

        k = len(sub) if len(sub) <= 4 else 3
        for r in range(k):
            row = [None, None, r + 1, to_excel_value(sub.loc[r, "Feature"]), to_excel_value(sub.loc[r, col])]
            if first_row_of_section:
                row[0] = title
                first_row_of_section = False
            if r == 0:
                row[1] = pair_label
            rows.append(row)

    return rows


# --------------------------
# Writer
# --------------------------
def write_summary_table(
    input_dir: Path,
    class_name_list: Sequence[str],
    number_of_attention_branches: int,
    model_type: str,  # "single" or "multi"
    output_path: Path,
) -> None:
    if len(class_name_list) < 2:
        raise ValueError("class_name_list must contain at least 2 condition names.")
    if number_of_attention_branches < 1:
        raise ValueError("number_of_attention_branches must be >= 1.")
    if model_type not in ("single", "multi"):
        raise ValueError("model_type must be 'single' or 'multi'.")

    sections = SECTIONS_SINGLE if model_type == "single" else SECTIONS_MULTI

    if model_type == "single":
        must = ["correlation.xlsx", "correlation_in_highlighted_segments.xlsx", "hellinger_distance.xlsx"]
    else:
        # Core files required in every multi run.
        must = [
            "correlation.xlsx",
            "correlation_in_highlighted_segments.xlsx",
            "hellinger_distance.xlsx",
            "correlation_node_diff.xlsx",
            "correlation_node_diff_in_highlighted_segments.xlsx",
            "hellinger_distance_node_diff.xlsx",
            "correlation_edge.xlsx",
            "correlation_edge_in_highlighted_segments.xlsx",
            "hellinger_distance_edge.xlsx",
            "correlation_edge_per_node.xlsx",
            "correlation_edge_per_node_in_highlighted_segments.xlsx",
            "hellinger_distance_edge_per_node.xlsx",
            "correlation_cluster.xlsx",
            "correlation_cluster_in_highlighted_segments.xlsx",
            "node_attention.xlsx",
        ]
        # MST files are optional (require a "Distance between nodes" edge feature).
        # Geometric files are optional (require x/y coordinates in feature_name_list).

    missing = [f for f in must if not (input_dir / f).exists()]
    if missing:
        raise FileNotFoundError(f"Missing required files in input_dir: {missing}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Attention sheets
    for att in range(1, number_of_attention_branches + 1):
        ws = wb.create_sheet(f"Attention{att}")

        all_rows: List[List[object]] = []
        for title, fname, kind in sections:
            if not (input_dir / fname).exists():
                # Skip optional sections (e.g., geometric features when coords are missing).
                print(f"[make_summary_table] Skipping missing file: {fname}")
                continue
            if kind == "corr":
                all_rows += build_corr_rows(input_dir, fname, att, title, class_name_list)
            elif kind == "edge":
                all_rows += build_edge_rows(input_dir, fname, att, title, class_name_list)
            elif kind == "cluster":
                all_rows += build_cluster_rows(input_dir, fname, att, title, class_name_list)
            elif kind == "hell":
                all_rows += build_hellinger_rows(input_dir, fname, att, title, class_name_list)
            else:
                raise ValueError(f"Unknown section kind: {kind}")

        for r_idx, row in enumerate(all_rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                ws.cell(r_idx, c_idx).value = val

    # Node attention: only for multi
    if model_type == "multi":
        node_path = input_dir / "node_attention.xlsx"
        validate_sheets_exist(node_path, class_name_list)

        ws = wb.create_sheet("Node attention")

        header = [None, None] + [f"Attention{i}" for i in range(1, number_of_attention_branches + 1)]
        for c, v in enumerate(header, start=1):
            ws.cell(1, c).value = v

        row_ptr = 2
        for cond in class_name_list:
            df = pd.read_excel(node_path, sheet_name=cond)
            node_col = "Unnamed: 0" if "Unnamed: 0" in df.columns else df.columns[0]

            # Attention columns check
            for i in range(1, number_of_attention_branches + 1):
                col = f"Attention{i}"
                if col not in df.columns:
                    raise ValueError(f"node_attention.xlsx:{cond} missing column '{col}'.")

            for _, r in df.iterrows():
                node_id = str(r[node_col])
                ws.cell(row_ptr, 1).value = cond
                ws.cell(row_ptr, 2).value = node_id
                for i in range(1, number_of_attention_branches + 1):
                    ws.cell(row_ptr, 2 + i).value = to_excel_value(r[f"Attention{i}"])
                row_ptr += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


# --------------------------
# CLI
# --------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input_dir", type=str, help="Folder containing the input xlsx files.")
    ap.add_argument(
        "--class-names",
        type=str,
        required=True,
        help="Python list literal (e.g. \"['a','b']\") OR comma-separated (e.g. \"a,b\").",
    )
    ap.add_argument("--num-attentions", type=int, required=True, help="number_of_attention_branches (e.g. 5)")
    ap.add_argument("--model-type", choices=["single", "multi"], required=True, help="model_type: single or multi")
    ap.add_argument("-o", "--output", type=str, default="SummaryTable.xlsx", help="Output xlsx path.")
    args = ap.parse_args()

    class_name_list = parse_class_names(args.class_names)

    write_summary_table(
        input_dir=Path(args.input_dir),
        class_name_list=class_name_list,
        number_of_attention_branches=args.num_attentions,
        model_type=args.model_type,
        output_path=Path(args.output),
    )


if __name__ == "__main__":
    main()
